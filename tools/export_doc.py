import os
import re
from tools.base import Tool

EXPORTS_DIR = "data/exports"


def _ensure_exports_dir() -> str:
    os.makedirs(EXPORTS_DIR, exist_ok=True)
    return EXPORTS_DIR


def _sanitize_filename(name: str, max_len: int = 50) -> str:
    """Remove unsafe characters from filename."""
    name = re.sub(r'[<>:"/\\|?*]', "", name).strip()
    return name[:max_len] if name else "export"


def _export_md(content: str, filepath: str) -> str:
    with open(filepath, "w", encoding="utf-8") as f:
        f.write(content)
    return filepath


def _export_txt(content: str, filepath: str) -> str:
    # Strip markdown formatting for plain text
    plain = content
    plain = re.sub(r"#{1,6}\s*", "", plain)          # headings
    plain = re.sub(r"\*\*(.+?)\*\*", r"\1", plain)    # bold
    plain = re.sub(r"\*(.+?)\*", r"\1", plain)        # italic
    plain = re.sub(r"`{1,3}[^`]*`{1,3}", "", plain)   # code
    plain = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", plain)  # links
    with open(filepath, "w", encoding="utf-8") as f:
        f.write(plain)
    return filepath


def _export_docx(content: str, filepath: str) -> str:
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    doc.styles["Normal"].font.size = Pt(11)

    for line in content.split("\n"):
        line = line.rstrip()
        if not line:
            doc.add_paragraph("")
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=2)
        elif line.startswith("# "):
            doc.add_heading(line[2:], level=1)
        elif line.startswith("### "):
            doc.add_heading(line[4:], level=3)
        elif line.startswith("- "):
            doc.add_paragraph(line[2:], style="List Bullet")
        elif re.match(r"^\d+\.\s", line):
            doc.add_paragraph(re.sub(r"^\d+\.\s", "", line), style="List Number")
        elif line.startswith("> "):
            doc.add_paragraph(line[2:], style="Quote")
        else:
            # Handle inline bold and italic
            p = doc.add_paragraph()
            parts = re.split(r"(\*\*.*?\*\*|\*.*?\*|`.*?`)", line)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    run = p.add_run(part[2:-2])
                    run.bold = True
                elif part.startswith("*") and part.endswith("*") and not part.startswith("**"):
                    run = p.add_run(part[1:-1])
                    run.italic = True
                elif part.startswith("`") and part.endswith("`"):
                    run = p.add_run(part[1:-1])
                    run.font.name = "Consolas"
                else:
                    p.add_run(part)

    doc.save(filepath)
    return filepath


def _export_html(content: str, filepath: str) -> str:
    # Simple markdown-to-HTML conversion using regex (no extra dependency)
    html = content
    html = re.sub(r"^### (.+)$", r"<h3>\1</h3>", html, flags=re.MULTILINE)
    html = re.sub(r"^## (.+)$", r"<h2>\1</h2>", html, flags=re.MULTILINE)
    html = re.sub(r"^# (.+)$", r"<h1>\1</h1>", html, flags=re.MULTILINE)
    html = re.sub(r"\*\*(.+?)\*\*", r"<strong>\1</strong>", html)
    html = re.sub(r"`([^`]+)`", r"<code>\1</code>", html)
    html = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r'<a href="\2">\1</a>', html)
    html = html.replace("\n", "<br>\n")

    full = f"""<!DOCTYPE html>
<html lang="zh-CN">
<head><meta charset="UTF-8"><title>Exported Document</title>
<style>
  body {{ font-family: -apple-system, sans-serif; max-width: 800px; margin: 40px auto; padding: 0 20px; line-height: 1.7; color: #333; }}
  h1 {{ color: #c0392b; border-bottom: 2px solid #c0392b; padding-bottom: 8px; }}
  h2 {{ color: #2c3e50; margin-top: 24px; }}
  h3 {{ color: #555; }}
  code {{ background: #f0f0f0; padding: 2px 6px; border-radius: 3px; font-size: 0.9em; }}
  strong {{ color: #c0392b; }}
  a {{ color: #2980b9; }}
</style></head>
<body>{html}</body></html>"""

    with open(filepath, "w", encoding="utf-8") as f:
        f.write(full)
    return filepath


EXPORTERS = {
    "md": ("Markdown", _export_md),
    "txt": ("Plain Text", _export_txt),
    "docx": ("Word Document", _export_docx),
    "html": ("HTML Page", _export_html),
}


class ExportDoc(Tool):
    name = "export_doc"
    description = (
        "Export content as a file in the specified format. "
        "Use this when the user wants to save research results, notes, or answers "
        "as a downloadable document. Supports: md (Markdown), txt (Plain Text), "
        "docx (Microsoft Word), html (Web page). Default is md."
    )
    parameters = {
        "type": "object",
        "properties": {
            "content": {
                "type": "string",
                "description": "The full content to export. Use Markdown formatting for best results across all formats."
            },
            "filename": {
                "type": "string",
                "description": "Base filename without extension, e.g. 'transformer-notes'. Default: auto-generated from content."
            },
            "format": {
                "type": "string",
                "enum": ["md", "txt", "docx", "html"],
                "description": "Output format: 'md' (Markdown, default), 'txt' (Plain Text), 'docx' (Word), 'html' (Web page)."
            }
        },
        "required": ["content"]
    }

    def execute(self, content: str, filename: str = "", format: str = "md", **kwargs) -> str:
        fmt = format.lower() if format else "md"
        if fmt not in EXPORTERS:
            return (
                f"Unsupported format: '{format}'. "
                f"Available formats: {', '.join(EXPORTERS.keys())}"
            )

        label, exporter_fn = EXPORTERS[fmt]

        if not filename:
            # Auto-generate filename from first line
            first_line = content.split("\n")[0].strip().lstrip("#").strip()
            filename = _sanitize_filename(first_line) if first_line else "export"

        filepath = os.path.join(_ensure_exports_dir(), f"{filename}.{fmt}")
        exporter_fn(content, filepath)

        return (
            f"Document exported successfully!\n"
            f"  Format: {label} (.{fmt})\n"
            f"  File: {filepath}\n"
            f"  Size: {os.path.getsize(filepath)} bytes"
        )
